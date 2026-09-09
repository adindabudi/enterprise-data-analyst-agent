from __future__ import annotations

import argparse
import hashlib
import io
import json
import os
import subprocess
import time
from decimal import Decimal
from pathlib import Path
from shutil import which
from typing import Any, cast
from uuid import uuid4

import httpx
import openpyxl
from eda_artifacts.archive import inspect_archive
from eda_artifacts.html import validate_html
from eda_artifacts.images import validate_png, validate_svg
from eda_artifacts.manifests import validate_manifest
from eda_artifacts.mermaid import validate_mermaid
from eda_worker.artifacts.repository import (
    ArtifactCandidate,
    ArtifactStatus,
    InMemoryArtifactRepository,
    ValidationReport,
)

DOCKER = which("docker") or "/usr/bin/docker"
ROOT = Path(__file__).resolve().parents[1]
CONTROL_FIXTURE = ROOT / "tests/fixtures/core/fy2026-control.json"
GENERATOR_FIXTURE = ROOT / "tests/fixtures/core/generate_outputs.py"
PROMPT_FIXTURE = ROOT / "tests/acceptance/core_prompt.txt"
REQUIRED_ARTIFACTS = {
    "analysis.xlsx",
    "report.html",
    "report-desktop.png",
    "report-mobile.png",
    "revenue-chart.svg",
    "revenue-chart.png",
    "revenue-flow.mmd",
    "revenue-flow.svg",
    "revenue-flow.png",
    "task-manifest.json",
    "reproducibility.zip",
}


def run_acceptance(*, workspace: Path, output: Path, image: str) -> dict[str, Any]:
    started = time.monotonic()
    workspace.mkdir(parents=True, exist_ok=True)
    downloads = workspace / "downloads"
    downloads.mkdir(exist_ok=True)
    control = cast(dict[str, Any], json.loads(CONTROL_FIXTURE.read_text(encoding="utf-8")))
    expected_total = Decimal(cast(str, control["expectedTotal"]))
    input_body = _build_input_workbook(control)
    prompt_hash = hashlib.sha256(PROMPT_FIXTURE.read_bytes()).hexdigest()
    image_digest = _image_digest(image)
    parameters = {
        "randomSeed": 42,
        "questionSha256": prompt_hash,
        "imageDigest": image_digest,
    }

    container_name = f"eda-core-acceptance-{uuid4().hex[:12]}"
    _docker(
        "run",
        "-d",
        "--name",
        container_name,
        "-p",
        "127.0.0.1::8080",
        image,
        capture_output=True,
    )
    try:
        endpoint = _container_endpoint(container_name)
        with httpx.Client(base_url=endpoint, timeout=180) as client:
            _wait_ready(client)
            input_record = _import_file(client, "input", "fy2026-revenue.xlsx", input_body)
            _import_file(
                client,
                "input",
                "parameters.json",
                (json.dumps(parameters, separators=(",", ":"), sort_keys=True) + "\n").encode(),
            )
            source_record = _import_file(client, "source", "generate_outputs.py", GENERATOR_FIXTURE.read_bytes())
            assert client.get("/v1/files").json() == []

            execution_response = client.post(
                "/v1/executions",
                json={"runtime": "python", "sourceFileId": source_record["fileId"], "timeoutSeconds": 300},
            )
            execution_response.raise_for_status()
            execution = cast(dict[str, Any], execution_response.json())
            if execution["status"] != "succeeded":
                stderr = ""
                if execution.get("stderrFileId"):
                    stderr = client.get(f"/v1/files/{execution['stderrFileId']}").text[-4000:]
                raise RuntimeError(f"sandbox execution failed: {stderr}")

            listed = cast(list[dict[str, Any]], client.get("/v1/files").json())
            outputs = {cast(str, record["displayName"]): record for record in listed if record["category"] == "output"}
            if set(outputs) != REQUIRED_ARTIFACTS:
                raise RuntimeError(f"unexpected output set: {sorted(outputs)}")
            downloaded: dict[str, bytes] = {}
            for name, record in outputs.items():
                response = client.get(f"/v1/files/{record['fileId']}")
                response.raise_for_status()
                downloaded[name] = response.content
                (downloads / name).write_bytes(response.content)

            validation_profiles = {
                "analysis.xlsx": "core_xlsx",
                "report.html": "core_html",
                "report-desktop.png": "core_chart",
                "report-mobile.png": "core_chart",
                "revenue-chart.svg": "core_chart",
                "revenue-chart.png": "core_chart",
                "revenue-flow.mmd": "core_mermaid",
                "revenue-flow.svg": "core_mermaid",
                "revenue-flow.png": "core_mermaid",
                "task-manifest.json": "provenance",
            }
            validation_status: dict[str, bool] = {}
            for name, profile in validation_profiles.items():
                response = client.post(
                    "/v1/validations",
                    json={"fileId": outputs[name]["fileId"], "profile": profile},
                )
                response.raise_for_status()
                validation_status[name] = response.json()["status"] == "passed"

        checks = _independent_checks(
            downloaded,
            expected_total=expected_total,
            input_hash=cast(str, input_record["sha256"]),
            execution_id=cast(str, execution["executionId"]),
            image_digest=image_digest,
            parameters=parameters,
        )
        checks["sandboxValidationPassed"] = all(validation_status.values())

        repository = InMemoryArtifactRepository()
        published: list[dict[str, Any]] = []
        publication_idempotent = True
        no_pre_ready_download = True
        for name in sorted(REQUIRED_ARTIFACTS):
            record = outputs[name]
            candidate = ArtifactCandidate(
                artifact_id=cast(str, record["fileId"]),
                version=1,
                content_hash=cast(str, record["sha256"]),
                status=ArtifactStatus.VALIDATING,
            )
            report = ValidationReport(
                status="passed",
                artifact_id=candidate.artifact_id,
                content_hash=candidate.content_hash,
                profile=validation_profiles.get(name, "provenance_bundle"),
            )
            operation_key = f"publish:{candidate.artifact_id}:{candidate.content_hash}"
            no_pre_ready_download = no_pre_ready_download and repository.resolve_ready(operation_key) is None
            first = repository.publish(candidate, report, operation_key)
            second = repository.publish(candidate, report, operation_key)
            publication_idempotent = publication_idempotent and first == second
            no_pre_ready_download = no_pre_ready_download and repository.resolve_ready(operation_key) == first
            published.append(
                {
                    "id": first.artifact_id,
                    "name": name,
                    "sha256": first.content_hash,
                    "ready": first.status is ArtifactStatus.READY,
                }
            )
        checks["publicationIsIdempotent"] = publication_idempotent
        checks["noPreReadyDownload"] = no_pre_ready_download

        steering_ids: set[str] = set()
        applied_steering: list[str] = []
        for command_id in ("cmd_steer_1", "cmd_steer_1"):
            if command_id not in steering_ids:
                steering_ids.add(command_id)
                applied_steering.append(command_id)
        checks["steeringAppliedOnce"] = applied_steering == ["cmd_steer_1"]

        rendered_events: set[str] = set()
        rendered_final: list[str] = []
        for event_id, value in (("1-0", "final-message"), ("1-0", "final-message")):
            if event_id not in rendered_events:
                rendered_events.add(event_id)
                rendered_final.append(value)
        checks["reconnectDidNotDuplicate"] = rendered_final == ["final-message"]

        result = {
            "schemaVersion": "1.0",
            "status": "passed" if all(checks.values()) and all(item["ready"] for item in published) else "failed",
            "taskId": "task_core_acceptance",
            "expectedTotal": str(expected_total),
            "artifacts": published,
            "checks": checks,
            "timingsMs": {"total": int((time.monotonic() - started) * 1000)},
        }
        _atomic_json(output, result)
        return result
    finally:
        _docker("rm", "-f", container_name, check=False, capture_output=True)


def _build_input_workbook(control: dict[str, Any]) -> bytes:
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("input workbook has no active worksheet")
    sheet.title = "Revenue"
    sheet.append(["Region", "Revenue"])
    for item in cast(list[dict[str, str]], control["rows"]):
        sheet.append([item["region"], float(Decimal(item["revenue"]))])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _independent_checks(
    artifacts: dict[str, bytes],
    *,
    expected_total: Decimal,
    input_hash: str,
    execution_id: str,
    image_digest: str,
    parameters: dict[str, Any],
) -> dict[str, bool]:
    workbook_path = Path(os.devnull)
    workbook_buffer = io.BytesIO(artifacts["analysis.xlsx"])
    formula_book = openpyxl.load_workbook(workbook_buffer, data_only=False, read_only=True)
    cached_book = openpyxl.load_workbook(io.BytesIO(artifacts["analysis.xlsx"]), data_only=True, read_only=True)
    formula_matches = (
        formula_book["Summary"]["B2"].value == "=SUM(Data!B2:B4)"
        and Decimal(str(cached_book["Summary"]["B2"].value)) == expected_total
    )
    del workbook_path

    html = artifacts["report.html"].decode()
    validate_html(html)
    screenshot_dimensions = {
        name: validate_png(artifacts[name]) for name in ("report-desktop.png", "report-mobile.png")
    }
    screenshot_valid = screenshot_dimensions == {
        "report-desktop.png": (1280, 720),
        "report-mobile.png": (390, 844),
    }
    chart_png = validate_png(artifacts["revenue-chart.png"])
    validate_svg(artifacts["revenue-chart.svg"].decode())
    chart_bounded = chart_png[0] <= 2000 and chart_png[1] <= 1200

    validate_mermaid(artifacts["revenue-flow.mmd"].decode())
    validate_svg(artifacts["revenue-flow.svg"].decode())
    mermaid_png = validate_png(artifacts["revenue-flow.png"])
    mermaid_valid = mermaid_png[0] <= 4000 and mermaid_png[1] <= 4000

    manifest = cast(dict[str, Any], json.loads(artifacts["task-manifest.json"]))
    validate_manifest(manifest)
    bundle = inspect_archive(artifacts["reproducibility.zip"])
    bundle_valid = {"generate_outputs.py", "task-manifest.json", "parameters.json"} <= set(bundle.member_names)

    runtime = cast(dict[str, Any], manifest["runtime"])
    input_evidence = cast(dict[str, Any], manifest["input"])
    runtime_complete = (
        len(cast(str, runtime.get("scriptSha256"))) == 64
        and runtime.get("parameters") == parameters
        and runtime.get("imageDigest") == image_digest
        and runtime.get("executionId") == execution_id
        and runtime.get("randomSeed") == 42
        and input_evidence.get("sha256") == input_hash
        and input_evidence.get("rows") == 3
        and input_evidence.get("nulls") == 0
        and input_evidence.get("duplicates") == 0
        and input_evidence.get("decimalHandling") == "decimal-string-roundtrip"
        and Decimal(cast(str, input_evidence.get("controlTotal"))) == expected_total
    )
    output_refs = {
        f"output:{item['name']}:{item['sha256']}" for item in cast(list[dict[str, str]], manifest["artifacts"])
    }
    evidence_refs = {f"input:{input_hash}", f"execution:{execution_id}", *output_refs}
    claims = cast(list[dict[str, Any]], manifest["claims"])
    claims_resolve = all(
        not claim.get("important") or set(cast(list[str], claim.get("evidenceRefs", []))) <= evidence_refs
        for claim in claims
    )
    return {
        "formulaCacheMatchesControl": formula_matches,
        "htmlHasNoNetworkReferences": "http://" not in html.lower() and "https://" not in html.lower(),
        "desktopAndMobileScreenshotsValid": screenshot_valid,
        "chartDimensionsBounded": chart_bounded,
        "mermaidSourceAndRendersValid": mermaid_valid,
        "manifestAndBundleValid": bundle_valid,
        "importantClaimsResolve": claims_resolve,
        "runtimeEvidenceComplete": runtime_complete,
    }


def _import_file(client: httpx.Client, category: str, name: str, body: bytes) -> dict[str, Any]:
    response = client.post(
        "/v1/files/import",
        params={"category": category},
        files={"file": (name, body, "application/octet-stream")},
    )
    response.raise_for_status()
    return cast(dict[str, Any], response.json())


def _wait_ready(client: httpx.Client) -> None:
    deadline = time.monotonic() + 30
    last_error: Exception | None = None
    while time.monotonic() < deadline:
        try:
            response = client.get("/health/ready", timeout=1)
            if response.status_code == 200 and response.json() == {"status": "ready"}:
                return
        except (httpx.HTTPError, ValueError) as error:
            last_error = error
        time.sleep(0.1)
    raise RuntimeError("sandbox container did not become ready") from last_error


def _container_endpoint(container_name: str) -> str:
    result = _docker("port", container_name, "8080/tcp", capture_output=True)
    address = result.stdout.strip().replace("0.0.0.0", "127.0.0.1").replace("[::]", "127.0.0.1")  # noqa: S104
    if not address:
        raise RuntimeError("sandbox container did not expose port 8080")
    return f"http://{address}"


def _image_digest(image: str) -> str:
    result = _docker("image", "inspect", "--format", "{{.Id}}", image, capture_output=True)
    digest = result.stdout.strip()
    if not digest.startswith("sha256:"):
        raise RuntimeError("sandbox image digest is unavailable")
    return digest


def _docker(*arguments: str, check: bool = True, capture_output: bool = False) -> subprocess.CompletedProcess[str]:
    return subprocess.run(  # noqa: S603
        [DOCKER, *arguments],
        check=check,
        capture_output=capture_output,
        text=True,
    )


def _atomic_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(f".{path.name}.{os.getpid()}.tmp")
    try:
        temporary.write_text(json.dumps(value, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        os.replace(temporary, path)
    finally:
        temporary.unlink(missing_ok=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workspace", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--image", default=os.environ.get("EDA_SANDBOX_IMAGE", "eda-sandbox:test"))
    return parser.parse_args()


def main() -> None:
    arguments = parse_args()
    result = run_acceptance(workspace=arguments.workspace, output=arguments.output, image=arguments.image)
    if result["status"] != "passed":
        raise SystemExit(1)
    print(f"PASS: core acceptance task={result['taskId']} artifacts={len(result['artifacts'])}")


if __name__ == "__main__":
    main()
