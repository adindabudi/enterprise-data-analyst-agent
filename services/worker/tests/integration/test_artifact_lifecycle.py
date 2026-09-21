from __future__ import annotations

import io
import shlex
import shutil
import subprocess
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest
from eda_contracts import ArtifactKind
from eda_runtime_state.models import RequiredOutput
from eda_runtime_state.tasks import InMemoryRuntimeStateRepository
from eda_worker.finalization import CoreTaskFinalizer
from eda_worker.history.repository import InMemoryProjectionRepository
from eda_worker.sandbox import aca_client
from eda_worker.sandbox.gateway import DynamicSessionCapabilityGateway, InMemoryArtifactGatewayStore
from eda_worker.tools.contracts import (
    CapabilityStatus,
    ExecuteSandboxOperation,
    PublishArtifactOperation,
    SandboxRuntime,
    ValidateArtifactOperation,
    ValidationProfile,
)
from openpyxl import load_workbook

from services.worker.tests.sandbox.test_gateway import _task


class LocalSandboxTransport:
    def __init__(self, root: Path) -> None:
        self.root = root
        self.deleted = False

    def write_file(self, path: str, content: bytes, **kwargs: object) -> None:
        Path(path).parent.mkdir(parents=True, exist_ok=True)
        Path(path).write_bytes(content)

    def read_file(self, path: str) -> bytes:
        return Path(path).read_bytes()

    def list_files(self, path: str) -> SimpleNamespace:
        return SimpleNamespace(
            entries=[
                SimpleNamespace(name=entry.name, path=str(entry), is_directory=entry.is_dir())
                for entry in Path(path).iterdir()
            ]
        )

    def exec(self, command: str, **kwargs: object) -> SimpleNamespace:
        if command.startswith("rm -rf -- "):
            shutil.rmtree(self.root / "outputs", ignore_errors=True)
            (self.root / "outputs").mkdir()
            return SimpleNamespace(stdout="", stderr="", exit_code=0)
        arguments = shlex.split(command)
        assert arguments[0] == "python"
        result = subprocess.run(  # noqa: S603
            [sys.executable, *arguments[1:]],
            cwd=self.root,
            capture_output=True,
            text=True,
            timeout=30,
            check=False,
        )
        return SimpleNamespace(stdout=result.stdout, stderr=result.stderr, exit_code=result.returncode)

    def delete(self) -> None:
        self.deleted = True


@pytest.mark.asyncio
@pytest.mark.parametrize("revenues", [(120, 150, 140), (17, 43, 91)])
async def test_real_execution_validation_publication_and_cleanup(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
    revenues: tuple[int, int, int],
) -> None:
    for name, value in {
        "WORKSPACE": str(tmp_path),
        "INPUTS": str(tmp_path / "inputs"),
        "SOURCES": str(tmp_path / "sources"),
        "OUTPUTS": str(tmp_path / "outputs"),
        "EXECUTION_CONTEXT": str(tmp_path / "execution-context.json"),
    }.items():
        monkeypatch.setattr(aca_client, name, value)
    transport = LocalSandboxTransport(tmp_path)
    client = aca_client.AcaSandboxClient(
        SimpleNamespace(get_sandbox_client=lambda identifier: transport),
        disk_image_id="disk_12345678",
    )
    runtime = InMemoryRuntimeStateRepository()
    store = InMemoryArtifactGatewayStore()
    task = _task()
    csv = "quarter,revenue\n" + "\n".join(f"Q{index},{amount}" for index, amount in enumerate(revenues, start=1))
    source_input = await store.persist_bytes(task.id, ArtifactKind.INPUT, "revenue.csv", csv.encode())
    task = task.model_copy(
        update={
            "input_artifacts": (source_input,),
            "active_sandbox_id": "sbx_12345678",
            "required_outputs": (RequiredOutput(kind=ArtifactKind.HTML), RequiredOutput(kind=ArtifactKind.XLSX)),
        }
    )
    await runtime.create_task(task, "request-lifecycle-12345678")
    gateway = DynamicSessionCapabilityGateway(client, runtime, store)
    finalizer = CoreTaskFinalizer(runtime, store, InMemoryProjectionRepository(), gateway)
    source = """import csv, json
from pathlib import Path
from openpyxl import Workbook
context = json.loads(Path("execution-context.json").read_text())
with Path(context["inputs"][0]["path"]).open() as handle:
    rows = list(csv.DictReader(handle))
total = sum(int(row["revenue"]) for row in rows)
output = Path(context["output_directory"])
workbook = Workbook()
workbook.active.append(["total", total])
workbook.save(output / "analysis.xlsx")
csp = "default-src 'none'; script-src 'unsafe-inline'; style-src 'unsafe-inline'; img-src data: blob:; font-src data:; connect-src 'none'; object-src 'none'; base-uri 'none'; form-action 'none'"
html = f'<html><head><meta http-equiv="Content-Security-Policy" content="{csp}"></head><body><h1>{total}</h1><script>document.body.dataset.ready="true";</script></body></html>'
(output / "dashboard.html").write_text(html)
print("generated from input")
"""
    result = await gateway.execute(
        task.id,
        ExecuteSandboxOperation(
            runtime=SandboxRuntime.PYTHON,
            source=source,
            input_artifacts=(source_input,),
            expected_outputs=("analysis.xlsx", "dashboard.html"),
        ),
    )
    assert result.status is CapabilityStatus.OK
    assert {artifact.kind for artifact in result.artifact_refs} == {ArtifactKind.HTML, ArtifactKind.XLSX}
    assert len(result.diagnostic_refs) == 2
    assert (await finalizer.validate_outputs({"taskId": task.id, "requiredProfile": "web_artifact_html"}))[
        "outcome"
    ] == "failed"
    for artifact in result.artifact_refs:
        payload = await store.read_bytes(task.id, artifact)
        if artifact.kind is ArtifactKind.XLSX:
            workbook = load_workbook(io.BytesIO(payload), data_only=True)
            assert workbook.active["B1"].value == sum(revenues)
            profile = ValidationProfile.CORE_XLSX
        else:
            assert f"<h1>{sum(revenues)}</h1>" in payload.decode()
            profile = ValidationProfile.WEB_ARTIFACT_HTML
        validation = await gateway.validate(task.id, ValidateArtifactOperation(artifact=artifact, profile=profile))
        assert validation.status is CapabilityStatus.OK
        publication = await gateway.publish(
            task.id,
            PublishArtifactOperation(
                artifact=artifact,
                validation_report=validation.artifact_refs[0],
            ),
        )
        assert publication.status is CapabilityStatus.OK
    assert (
        await finalizer.validate_outputs(
            {
                "taskId": task.id,
                "requiredProfile": "web_artifact_html",
                "requireOutputContract": True,
            }
        )
    )["outcome"] == "passed"
    await finalizer.complete_chat({"taskId": task.id, "text": "Verified results are ready."})
    assert transport.deleted
    restored = await runtime.resolve_task(task.id)
    assert restored is not None and restored.active_sandbox_id is None
