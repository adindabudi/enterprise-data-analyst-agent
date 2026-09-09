"""The acceptance run proves each document kind is generated deterministically.

Document Pack readiness compares the digest of two independent generations of
every kind. A generator that embeds a timestamp, or orders its archive by
dictionary iteration, passes locally and then fails acceptance in the cloud with
nothing to point at.
"""

from __future__ import annotations

import io
import zipfile

from eda_artifacts.documents import (
    ZIP_TIMESTAMP,
    generate_document_corpus,
    generate_xlsx,
    validate_generated_document,
)
from eda_artifacts.office import validate_office
from openpyxl import load_workbook

TITLE = "Enterprise Data Analyst"


def test_every_kind_the_pack_promises_is_generated() -> None:
    corpus = generate_document_corpus(TITLE)

    assert {document.kind for document in corpus} == {"docx", "pdf", "pptx", "xlsx"}


def test_generating_twice_produces_identical_bytes() -> None:
    first, second = generate_document_corpus(TITLE), generate_document_corpus(TITLE)

    for left, right in zip(first, second, strict=True):
        assert left.content == right.content, f"{left.kind} is not deterministic"


def test_every_archive_member_carries_the_fixed_timestamp_in_sorted_order() -> None:
    # Two generations one second apart in one process would agree anyway; acceptance compares
    # separate runs, so the properties that make that hold are asserted directly.
    for document in generate_document_corpus(TITLE):
        if not document.content.startswith(b"PK"):
            continue
        with zipfile.ZipFile(io.BytesIO(document.content)) as archive:
            members = archive.infolist()
            assert [member.filename for member in members] == sorted(member.filename for member in members)
            assert {member.date_time for member in members} == {ZIP_TIMESTAMP}


def test_each_generated_document_passes_its_own_validator() -> None:
    for document in generate_document_corpus(TITLE):
        validate_generated_document(document)


def test_the_workbook_is_a_valid_office_package() -> None:
    report = validate_office(generate_xlsx(TITLE), kind="xlsx")

    assert report.passed, report.failures


def test_a_real_spreadsheet_reader_opens_the_workbook() -> None:
    # Our own validator only inspects the package; this proves the file is actually readable.
    book = load_workbook(io.BytesIO(generate_xlsx(TITLE)))

    assert book.sheetnames == ["Analysis"]
    assert book["Analysis"]["A1"].value == TITLE


def test_a_title_containing_markup_cannot_break_the_workbook() -> None:
    book = load_workbook(io.BytesIO(generate_xlsx("<script>&amp;</script>")))

    assert book["Analysis"]["A1"].value == "<script>&amp;</script>"
